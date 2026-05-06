#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
IMP Excel报表生成器
扫描 /DATA/data2/downloads/IMP/ 各物种目录，生成xlsx数据报表
"""

import os
import argparse
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

BASE_DIR = "/home/nizhu/Projects/plantsdb"
DOWNLOAD_DIR = "/DATA/data2/downloads/IMP"

# 文件类型配置
FILE_TYPES = {
    "genome": {
        "extensions": [".fna", ".fa", ".fasta", "_genome.fna", "_genome.fa"],
        "column": "Genome_Sequence",
    },
    "annotation": {
        "extensions": [".gff", ".gff3", ".gff.gz", "_annotation.gff", "_genomic.gff"],
        "column": "Genome_Annotation",
    },
    "gene": {
        "extensions": [".fna", ".fa", "_gene.fna", "_genes.fna"],
        "column": "Gene_Sequence",
    },
    "cds": {
        "extensions": [".fna", "_cds.fna", "_cds_from_genomic.fna"],
        "column": "CDS_Sequence",
    },
    "protein": {
        "extensions": [".faa", ".fa", "_protein.faa", "_protein.fa"],
        "column": "Protein_Sequence",
    },
    "promoter": {
        "extensions": [".fna", ".fa", "_promoter.fna", "_promoter.fa"],
        "column": "Promoter_Sequence",
    },
    "tpm": {
        "extensions": [".tsv", ".txt", ".matrix", "_TPM.tsv", "_expression.tsv"],
        "column": "Gene_Expression_TPM",
    },
}


def check_file_exists(species_dir, file_type, species_name=None):
    """检查特定文件类型的文件是否存在"""
    if not os.path.exists(species_dir):
        return False

    # 根据物种全称查找可能的文件名
    patterns = {
        "genome": [f"{species_name}_genome.fa.gz", f"{species_name}_genome.fa", f"{species_name}.fa.gz"] if species_name else [],
        "annotation": [f"{species_name}_annotation.gff3.gz", f"{species_name}.gff3.gz"] if species_name else [],
        "gene": [f"{species_name}_gene.fasta"],
        "cds": [f"{species_name}_cds.fasta"],
        "protein": [f"{species_name}_protein.fasta"],
        "promoter": [f"{species_name}_promoter.fasta"],
        "tpm": [f"{species_name}_expression_TPM.txt"],
    }

    expected_names = patterns.get(file_type, [])

    for filename in os.listdir(species_dir):
        filepath = os.path.join(species_dir, filename)
        if not os.path.isfile(filepath):
            continue

        # 如果有明确的名字模式，精确匹配
        if expected_names and filename in expected_names:
            return True

        # 否则按扩展名模糊匹配（兼容性）
        if file_type == "genome" and ('genome' in filename.lower() or filename.endswith('.fa.gz')):
            return True
        elif file_type == "annotation" and ('annotation' in filename.lower() or filename.endswith('.gff3.gz')):
            return True
        elif file_type == "gene" and 'gene' in filename.lower() and filename.endswith('.fasta'):
            return True
        elif file_type == "cds" and 'cds' in filename.lower() and filename.endswith('.fasta'):
            return True
        elif file_type == "protein" and 'protein' in filename.lower() and filename.endswith('.fasta'):
            return True
        elif file_type == "promoter" and 'promoter' in filename.lower():
            return True
        elif file_type == "tpm" and ('tpm' in filename.lower() or 'expression' in filename.lower()):
            return True

    return False


def scan_species_directory(base_dir):
    """扫描下载目录，返回所有物种及其文件状态"""
    species_data = []

    if not os.path.exists(base_dir):
        return species_data

    # 读取species_list.json获取物种名称映射
    species_map = {}
    json_file = os.path.join(base_dir, "species_list.json")
    if os.path.exists(json_file):
        import json
        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
            for sp in data.get('species', []):
                code = sp.get('code', '')
                name = sp.get('name', code)
                dir_name = sp.get('dir', name.replace(' ', '_'))
                species_map[dir_name] = {'code': code, 'name': name}

    for folder_name in os.listdir(base_dir):
        folder_path = os.path.join(base_dir, folder_name)
        if not os.path.isdir(folder_path) or folder_name.startswith("."):
            continue

        # 跳过特殊目录
        if folder_name in ["logs", "static"]:
            continue

        # 获取物种信息
        sp_info = species_map.get(folder_name, {})
        species_code = sp_info.get('code', folder_name)
        species_name = sp_info.get('name', folder_name)

        # 检查各类文件
        file_status = {}
        for file_type, config in FILE_TYPES.items():
            file_status[config["column"]] = check_file_exists(folder_path, file_type, species_name)

        file_status["Species"] = species_name
        file_status["Source"] = "IMP"

        species_data.append(file_status)

    return species_data


def write_excel(data, output_file):
    """将数据写入Excel文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IMP Data Inventory"

    # 定义样式
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    cell_alignment = Alignment(horizontal="center", vertical="center")

    # 表头
    headers = ["Species", "Genome_Sequence", "Genome_Annotation",
               "Gene_Sequence", "CDS_Sequence", "Protein_Sequence",
               "Promoter_Sequence", "Gene_Expression_TPM", "Source"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # 数据行
    yes_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    no_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, "")

            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = cell_alignment

            if header == "Species":
                cell.value = value
            elif header == "Source":
                cell.value = "IMP"
            else:
                # 布尔值转为1/0
                cell.value = 1 if value else 0
                if value:
                    cell.fill = yes_fill
                else:
                    cell.fill = no_fill

    # 设置列宽
    column_widths = {
        "A": 25,  # Species
        "B": 15,  # Genome_Sequence
        "C": 18,  # Genome_Annotation
        "D": 15,  # Gene_Sequence
        "E": 15,  # CDS_Sequence
        "F": 16,  # Protein_Sequence
        "G": 18,  # Promoter_Sequence
        "H": 20,  # Gene_Expression_TPM
        "I": 10,  # Source
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 添加统计行
    last_row = len(data) + 2
    ws.cell(row=last_row, column=1, value="统计").font = Font(bold=True)

    for col_idx, header in enumerate(headers, 1):
        if header not in ["Species", "Source"]:
            count = sum(1 for row in data if row.get(header, False))
            ws.cell(row=last_row, column=col_idx, value=count)

    # 保存
    wb.save(output_file)


def main():
    parser = argparse.ArgumentParser(description="IMP Excel报表生成器")
    parser.add_argument("--indir", type=str, default=DOWNLOAD_DIR, help="输入目录")
    parser.add_argument("--out", type=str, default=None, help="输出xlsx文件路径")
    args = parser.parse_args()

    if not args.out:
        args.out = os.path.join(args.indir, "IMP_data_inventory.xlsx")

    print("=" * 60)
    print("IMP Excel报表生成器")
    print("=" * 60)
    print(f"输入目录: {args.indir}")
    print(f"输出文件: {args.out}")

    # 扫描目录
    print("\n扫描物种目录...")
    species_data = scan_species_directory(args.indir)
    print(f"找到物种数: {len(species_data)}")

    if not species_data:
        print("❌ 未找到任何物种数据")
        return

    # 写入Excel
    print("\n生成Excel报表...")
    write_excel(species_data, args.out)
    print(f"✅ 报表已生成: {args.out}")

    # 打印统计
    print("\n文件类型统计:")
    for file_type, config in FILE_TYPES.items():
        col = config["column"]
        count = sum(1 for row in species_data if row.get(col, False))
        print(f"  {col}: {count}/{len(species_data)}")


if __name__ == "__main__":
    main()
